import sqlite3
from openpyxl import *

class table():
    def __init__(self):
        pass

    pass

class excel_table(table):

    def ins_into_excel(self, work_info, db_path):
        wb = load_workbook(db_path)
        ws = wb.active
        work_info = [self.excel_safe_value(el) for el in work_info]
        ws.append(work_info)

        wb.save(db_path)

        wb.close()
        pass

    def excel_safe_value(self, value):
        """Преобразует значение для безопасной записи в Excel"""
        if value is None:
            return ""

        if isinstance(value, (str, bool)):
            return value

        if isinstance(value, (int, float)):
            if abs(value) > 1e15:
                return f"'{value}"
            return value

        if isinstance(value, (dict, list)):
            return str(value)

        return str(value)

    pass


class SQL_db(table):

    def create_sql_table(self, path):
        db = sqlite3.connect(path)
        curs = db.cursor()
        curs.execute("""
            CREATE TABLE resTable(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            work_id INTEGER,
            E1Numb1 integer,
            E1Numb2 integer, 
            E1Numb3 integer, 
            E1Numb4 integer, 
            E1Numb5 integer, 
            E1Numb6 integer,
            E2Numb1 integer,
            E2Numb2 integer, 
            E2Numb3 integer, 
            E2Numb4 integer, 
            E2Numb5 integer, 
            E2Numb6 integer,   
            id_exp1 integer,
            id_exp2 integer );""")
        db.commit()
        db.close()
        pass

    def create_exell_tab(self, db_path: str):
        db = sqlite3.connect('db_dir/ege_works.db')
        res_table_headers = ["ID", "WORK_ID", "E1 N1", "E1 N2", "E1 N3", "E1 N4", "E1 N5", "E1 N6",
                             "E2 N1", "E2 N2", "E2 N3", "E2 N4", "E2 N5", "E2 N6", 'id_exp1', 'id_exp2']
        wb = Workbook()
        ws = wb.active
        for name, col in enumerate("ABCDEFGHIJKLMNOP", start=0):
            ws[f'{col}1'] = res_table_headers[name]
        wb.save(db_path)
        db.close()
        pass

    def get_line(self, line_numb: int, sql_db_path: str):
        db = sqlite3.connect(sql_db_path)
        curs = db.cursor()

        curs.execute(f"SELECT * FROM resTable WHERE id = ?", (line_numb,))

        raw = curs.fetchone()
        print(raw)
        return raw

    def ins_into_sql(self, work_id, rates):
        db = sqlite3.connect('db_dir/ege_works.db')
        curs = db.cursor()

        curs.execute(f"""INSERT INTO resTable (work_id, E1Numb1, E1Numb2,E1Numb3,E1Numb4,E1Numb5,E1Numb6,E2Numb1,E2Numb2,
        E2Numb3,E2Numb4,E2Numb5,E2Numb6, id_exp1, id_exp2) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?) """,
                     (work_id,
                      rates['exp1'][0], rates['exp1'][1], rates['exp1'][2], rates['exp1'][3], rates['exp1'][4],
                      rates['exp1'][5],
                      rates['exp2'][0], rates['exp2'][1], rates['exp2'][2], rates['exp2'][3], rates['exp2'][4],
                      rates['exp2'][5],
                      rates['id_exp1'], rates['id_exp2'])
                     )
        db.commit()
        db.close()
        pass




if __name__ == "__main__":
    import os

    user_excel_table = excel_table()
    if not os.path.exists(excel_db_path):
        user_excel_table.create_exell_tab(excel_db_path) # user table
    create_sql_table('db_dir/ege_works.db')
