"""Validate recognized pages, compare two experts and export results safely."""

from contextlib import closing, contextmanager
import json
from pathlib import Path
import sqlite3
import tempfile


SCORE_COUNT = 6
VALID_SCORES = {"0", "1", "2", "3", "4", "X"}
HEADERS = (
    ["WORK_ID"]
    + [f"E{expert} N{n}" for expert in (1, 2) for n in range(1, SCORE_COUNT + 1)]
    + ["id_exp1", "id_exp2", "mistake", "n_w_mistake"]
)


def identifier(value, label):
    # IDs stay strings: Excel and SQLite must preserve leading zeros.
    if not isinstance(value, str) or not value or not value.isascii() or not value.isdigit():
        raise ValueError(f"{label} must be a non-empty string of ASCII digits; got {value!r}")
    return value


def compare_scores(first, second):
    if len(first) != SCORE_COUNT or len(second) != SCORE_COUNT:
        raise ValueError("Each expert must supply exactly six scores")
    differences = []
    normalized = []
    for scores in (first, second):
        values = [str(value).upper() for value in scores]
        if any(value not in VALID_SCORES for value in values):
            raise ValueError("Scores must be 0, 1, 2, 3, 4 or X")
        normalized.append(values)
    for n, (a, b) in enumerate(zip(*normalized), start=1):
        if (a != b if "X" in (a, b) else abs(int(a) - int(b)) >= 2):
            differences.append(n)
    return normalized, differences


def build_results(pages):
    if not isinstance(pages, list) or not pages:
        raise ValueError("Expected a non-empty list of recognized pages")
    works = {}
    for page in pages:
        if not isinstance(page, dict):
            raise ValueError("Each page must contain expert_id and works")
        expert = identifier(page.get("expert_id"), "expert_id")
        page_works = page.get("works")
        if not isinstance(page_works, dict) or not page_works:
            raise ValueError("Each page must contain a non-empty works object")
        for work_id, scores in page_works.items():
            identifier(work_id, "work_id")
            if not isinstance(scores, list):
                raise ValueError(f"Scores for work {work_id} must be a list")
            compare_scores(scores, scores)
            experts = works.setdefault(work_id, {})
            if expert in experts:
                raise ValueError(f"Duplicate expert {expert} for work {work_id}")
            experts[expert] = scores
    rows = []
    for work_id, experts in sorted(works.items()):
        if len(experts) != 2:
            raise ValueError(f"Work {work_id}: expected exactly two different experts, got {len(experts)}")
        ids = sorted(experts)
        scores, differences = compare_scores(experts[ids[0]], experts[ids[1]])
        values = [int(value) if value != "X" else value for value in scores[0] + scores[1]]
        rows.append([work_id, *values, *ids, len(differences), " ".join(map(str, differences))])
    return rows


@contextmanager
def staged_output(output_dir):
    """Build a new result directory without touching existing outputs."""
    destination = Path(output_dir).resolve()
    if destination.exists():
        raise FileExistsError(f"Output directory already exists: {destination}. Choose a new directory.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".ratescaner-", dir=destination.parent) as temporary:
        staged = Path(temporary) / "result"
        staged.mkdir()
        yield staged, destination
        staged.rename(destination)


def write_pages(pages, directory):
    (directory / "recognized_pages.json").write_text(
        json.dumps(pages, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def process_pages(pages, output_dir):
    """Keep recognized data for correction, even if semantic validation fails."""
    try:
        rows = build_results(pages)
    except ValueError as error:
        with staged_output(output_dir) as (staged, destination):
            write_pages(pages, staged)
            (staged / "review_required.txt").write_text(str(error) + "\n", encoding="utf-8")
        raise ValueError(
            f"{error}. Data saved to {destination / 'recognized_pages.json'}; "
            "correct it and rerun --from-json with a new --output directory."
        ) from error
    return export_results(rows, output_dir, recognized_pages=pages)


def export_results(rows, output_dir, *, recognized_pages=None):
    """Publish exports and optional source data in a new directory."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    with staged_output(output_dir) as (staged, destination):
        if recognized_pages is not None:
            write_pages(recognized_pages, staged)
        with closing(sqlite3.connect(staged / "results.sqlite3")) as connection:
            columns = (
                ['"WORK_ID" TEXT PRIMARY KEY']
                + [f'"E{expert} N{n}"' for expert in (1, 2) for n in range(1, SCORE_COUNT + 1)]
                + ['id_exp1 TEXT', 'id_exp2 TEXT', 'mistake INTEGER', 'n_w_mistake TEXT']
            )
            connection.execute(f"CREATE TABLE results ({', '.join(columns)})")
            connection.executemany("INSERT INTO results VALUES (" + ",".join(["?"] * len(HEADERS)) + ")", rows)
            connection.commit()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "comparison"
        sheet.append(HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17324D")
        for row in rows:
            sheet.append(row)
            for position in row[-1].split():
                for column in (1 + int(position), 1 + SCORE_COUNT + int(position)):
                    sheet.cell(sheet.max_row, column).fill = PatternFill("solid", fgColor="FFF0BF")
        for column in range(1, len(HEADERS) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 9 if 2 <= column <= 13 else 22
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(staged / "results.xlsx")
        workbook.close()
    return destination
