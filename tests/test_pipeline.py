import copy
import json
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import unittest

from scanner_pipeline import build_results, export_results

ROOT = Path(__file__).resolve().parents[1]


def pages():
    return json.loads((ROOT / "examples/recognized_pages.json").read_text())


class ComparisonTests(unittest.TestCase):
    def test_numeric_and_x_differences_and_leading_zeros(self):
        rows = build_results(pages())
        self.assertEqual(rows[0][0], "000001")
        self.assertEqual(rows[0][13:15], ["001", "002"])
        self.assertEqual(rows[0][-2:], [2, "2 4"])
        self.assertEqual(rows[1][-2:], [0, ""])

    def test_page_order_does_not_change_results(self):
        self.assertEqual(build_results(pages()), build_results(list(reversed(pages()))))

    def test_missing_or_duplicate_expert_is_rejected(self):
        data = pages()
        for candidate in (data[:1], [data[0], copy.deepcopy(data[0])]):
            with self.subTest(candidate=candidate), self.assertRaises(ValueError):
                build_results(candidate)

    def test_third_expert_is_not_silently_overwritten(self):
        data = pages()
        extra = copy.deepcopy(data[0])
        extra["expert_id"] = "003"
        with self.assertRaisesRegex(ValueError, "exactly two"):
            build_results(data + [extra])

    def test_invalid_scores_and_identifiers_are_rejected(self):
        for value in ("", "?", 5, -1, None):
            data = pages()
            data[0]["works"]["000001"][0] = value
            with self.subTest(value=value), self.assertRaises(ValueError):
                build_results(data)
        for expert in ("", "abc", "１２３", 123, None):
            data = pages()
            data[0]["expert_id"] = expert
            with self.subTest(expert=expert), self.assertRaises(ValueError):
                build_results(data)

    def test_incomplete_score_list_is_rejected(self):
        data = pages()
        data[0]["works"]["000001"].pop()
        with self.assertRaisesRegex(ValueError, "six scores"):
            build_results(data)


class ExportTests(unittest.TestCase):
    def test_sqlite_and_excel_have_expected_content(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "result"
            rows = build_results(pages())
            export_results(rows, output)
            with sqlite3.connect(output / "results.sqlite3") as connection:
                actual = connection.execute('SELECT WORK_ID, id_exp1, mistake, n_w_mistake FROM results ORDER BY WORK_ID').fetchall()
            self.assertEqual(actual, [("000001", "001", 2, "2 4"), ("000002", "001", 0, "")])
            workbook = load_workbook(output / "results.xlsx")
            sheet = workbook.active
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(sheet["A2"].value, "000001")
            self.assertEqual(sheet["N2"].value, "001")
            self.assertEqual(sheet["P2"].value, 2)
            self.assertEqual(sheet["Q2"].value, "2 4")
            workbook.close()

    def test_existing_output_is_preserved(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "result"
            output.mkdir()
            marker = output / "keep.txt"
            marker.write_text("original")
            with self.assertRaises(FileExistsError):
                export_results(build_results(pages()), output)
            self.assertEqual(marker.read_text(), "original")
            self.assertEqual(list(output.iterdir()), [marker])

    def test_cli_runs_from_another_directory(self):
        with tempfile.TemporaryDirectory() as temporary:
            result = subprocess.run(
                [sys.executable, str(ROOT / "main.py"), "--from-json",
                 str(ROOT / "examples/recognized_pages.json"), "--output", str(Path(temporary) / "out")],
                cwd=temporary, capture_output=True, text=True, timeout=20,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue((Path(temporary) / "out/results.xlsx").is_file())


if __name__ == "__main__":
    unittest.main()
