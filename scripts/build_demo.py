"""Build the downloadable demo workbook and Markdown preview from checked pages.

Run from the repository root: python scripts/build_demo.py
This does not run OCR; ocr_pages.json is a separately captured model output.
"""

import json
from pathlib import Path
import shutil
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from scanner_pipeline import process_pages


def main():
    demo = ROOT / "examples/protocol_demo"
    expected = json.loads((demo / "expected_pages.json").read_text(encoding="utf-8"))
    actual = json.loads((demo / "ocr_pages.json").read_text(encoding="utf-8"))
    preview = []
    with tempfile.TemporaryDirectory() as temporary:
        result = process_pages(expected, Path(temporary) / "corrected")
        shutil.copyfile(result / "results.xlsx", demo / "results.xlsx")
        workbook = load_workbook(result / "results.xlsx", read_only=True)
        for row in workbook.active.iter_rows(min_row=2, values_only=True):
            preview.append(
                f"| `{row[0]}` | {' '.join(map(str, row[1:7]))} | "
                f"{' '.join(map(str, row[7:13]))} | {row[15]} | {row[16] or '—'} |"
            )
        workbook.close()
    mismatches = []
    actual_by_expert = {page["expert_id"]: page["works"] for page in actual}
    if set(actual_by_expert) != {page["expert_id"] for page in expected}:
        raise ValueError("Expert identifiers differ; review the OCR snapshot")
    for page_number, page in enumerate(expected, 1):
        actual_works = actual_by_expert[page["expert_id"]]
        if set(actual_works) != set(page["works"]):
            raise ValueError("Work identifiers differ; review the OCR snapshot")
        for work, scores in page["works"].items():
            recognized = actual_works[work]
            if len(recognized) != len(scores):
                raise ValueError("Score counts differ; review the OCR snapshot")
            for position, (found, correct) in enumerate(zip(recognized, scores), 1):
                if str(found) != str(correct):
                    mismatches.append(f"| {page_number} | `{work}` | {position} | {found} | {correct} |")
    text = """# Демонстрация: бланки → OCR → проверка → Excel

Это два созданных автором демонстрационных протокола, а не реальные данные экзамена.
Читаются первые пять работ каждого бланка (по шесть оценок). Позиции 1–6 в отчёте
соответствуют полям 17–22 исходного бланка.

## Исходные изображения

[Бланк 1](../scans/1.jpg) · [Бланк 2](../scans/2.jpg)

<img src="../scans/1.jpg" alt="Первый демонстрационный протокол оценивания" width="760">

## Данные и результат

- [Исходный снимок OCR](../examples/protocol_demo/ocr_pages.json) — результат модели
  и Tesseract без ручного исправления.
- [Проверенная расшифровка](../examples/protocol_demo/expected_pages.json) — оценки
  и идентификаторы, сверенные по двум изображениям. Это ручная расшифровка, а не
  утверждение, что модель распознала их правильно.
- [Скачать Excel после проверки](../examples/protocol_demo/results.xlsx).

Ниже сводка из сформированной Excel-книги. В ней оранжевым выделены пары оценок
с расхождениями: разница от двух баллов либо несовпадение X. Программа не определяет,
какой эксперт прав. E1/E2 упорядочены по ID эксперта, а не по номеру изображения:
E1 — 133722813124940 (бланк 2), E2 — 1488902412421 (бланк 1).

| Работа | Оценки E1 | Оценки E2 | Расхождений | Позиции |
| --- | --- | --- | --- | --- |
""" + "\n".join(preview) + """

## Что потребовало ручной коррекции

В сохранённом OCR-снимке найдены следующие несовпадения с ручной расшифровкой:

| Бланк | Работа | Позиция | OCR | На изображении |
| --- | --- | --- | --- | --- |
""" + "\n".join(mismatches) + """

Этот маленький набор не измеряет общую точность и не является независимым бенчмарком.
Вес и окружение OCR-снимка описаны в [протоколе проверки](verification.md).

## Повторить без модели и Tesseract

Из корня репозитория, в окружении с openpyxl:

```bash
python main.py --from-json examples/protocol_demo/expected_pages.json --output outputs/demo-reviewed
```

Для сравнения с неисправленным OCR можно отдельно экспортировать снимок:

```bash
python main.py --from-json examples/protocol_demo/ocr_pages.json --output outputs/demo-raw
```

Оба каталога должны быть новыми. Экспорт снимка не означает, что OCR проверен.

Для пересборки приложенного Excel и этой страницы: `python scripts/build_demo.py`.
Скрипт использует сохранённые JSON, не запускает распознавание и не переписывает их.
Он обновляет только демонстрационный Excel и эту страницу.
"""
    (ROOT / "docs/demo.md").write_text(text, encoding="utf-8")
    print(f"Demo workbook and preview built; OCR mismatches: {len(mismatches)}")


if __name__ == "__main__":
    main()
